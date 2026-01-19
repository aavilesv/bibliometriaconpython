from docx import Document

doc = Document()
doc.add_heading('Subtemas para Antecedentes y Marco Teórico', level=1)

doc.add_paragraph(
    "Este documento organiza los subtemas sugeridos para la investigación "
    "“Formación de competencias lingüísticas de los estudiantes en espacios de enseñanza de la carrera de Educación "
    "en una universidad estatal del Ecuador, según las voces de los docentes y discentes”, "
    "diferenciando claramente los ítems correspondientes a los Antecedentes de investigación y al Marco Teórico. "
    "Cada subtema incluye su objetivo, palabras clave y ecuaciones de búsqueda para Scopus y Web of Science."
)

# Helper function
def add_section(title, objective, keywords, equations):
    doc.add_heading(title, level=2)
    doc.add_paragraph("Objetivo del subtema:\n" + objective)
    doc.add_paragraph("Palabras clave y sinónimos:")
    for k in keywords:
        doc.add_paragraph(k, style='List Bullet')
    doc.add_paragraph("Ecuaciones de búsqueda sugeridas (Scopus / Web of Science):")
    for e in equations:
        doc.add_paragraph(e, style='List Number')

doc.add_heading('I. Subtemas para Antecedentes de Investigación', level=1)

add_section(
    "A1. Formación de competencias lingüísticas en estudiantes universitarios de Educación",
    "Analizar estudios empíricos que abordan la formación de competencias lingüísticas en estudiantes universitarios de la carrera de Educación, identificando enfoques metodológicos y principales hallazgos.",
    [
        "Linguistic competence", "Language competence", "Academic language",
        "Teacher education students", "Pre-service teachers", "Higher education"
    ],
    [
        '("linguistic competence" OR "language competence" OR "academic language") AND ("teacher education students" OR "pre-service teachers" OR "education degree") AND ("higher education")'
    ]
)

add_section(
    "A2. Prácticas docentes y desarrollo de competencias lingüísticas en la educación superior",
    "Identificar investigaciones que analicen la relación entre prácticas docentes universitarias y el desarrollo de competencias lingüísticas en estudiantes.",
    [
        "Teaching practices", "Pedagogical practices", "Instructional practices",
        "Language development", "Linguistic skills"
    ],
    [
        '("teaching practices" OR "pedagogical practices") AND ("language development" OR "linguistic skills") AND ("higher education")'
    ]
)

add_section(
    "A3. Percepciones y creencias de docentes y estudiantes sobre la formación lingüística",
    "Examinar estudios que exploren las percepciones, creencias o representaciones de docentes y estudiantes respecto a la formación lingüística en la educación superior.",
    [
        "Teacher beliefs", "Student perceptions", "Student voice",
        "Social representations", "Language learning beliefs"
    ],
    [
        '("teacher beliefs" OR "student perceptions" OR "student voice") AND ("language learning" OR "linguistic competence") AND ("higher education")'
    ]
)

add_section(
    "A4. Alfabetización académica en la formación docente",
    "Revisar investigaciones empíricas sobre alfabetización académica y escritura académica en estudiantes de formación docente.",
    [
        "Academic literacy", "Academic writing", "Disciplinary literacy",
        "Teacher training", "Higher education"
    ],
    [
        '("academic literacy" OR "academic writing") AND ("teacher education" OR "teacher training") AND ("higher education")'
    ]
)

add_section(
    "A5. Estudios cualitativos sobre lenguaje y formación docente",
    "Identificar estudios cualitativos o mixtos que analicen el lenguaje, el discurso y las prácticas comunicativas en la formación inicial docente.",
    [
        "Qualitative study", "Mixed methods", "Discourse analysis",
        "Language practices", "Teacher education"
    ],
    [
        '("teacher education" AND "language") AND ("qualitative study" OR "mixed methods")'
    ]
)

doc.add_heading('II. Subtemas para el Marco Teórico', level=1)

add_section(
    "T1. Prácticas docentes mediadoras del lenguaje",
    "Analizar el rol de las prácticas docentes como mediadoras del lenguaje en los procesos de enseñanza y aprendizaje en la educación superior.",
    [
        "Teaching practices", "Pedagogical mediation", "Language mediation",
        "Classroom discourse", "Teacher talk"
    ],
    [
        '("pedagogical mediation" OR "language mediation") AND ("teaching practices") AND ("higher education")'
    ]
)

add_section(
    "T2. Competencias lingüísticas como componente del perfil profesional docente",
    "Describir las competencias lingüísticas como parte del perfil profesional del futuro docente y su relevancia en la mediación pedagógica.",
    [
        "Linguistic competence", "Professional competence", "Teacher profile",
        "Academic language", "Teacher education"
    ],
    [
        '("linguistic competence" AND "professional competence") AND ("teacher education" OR "initial teacher training")'
    ]
)

add_section(
    "T3. Espacios de enseñanza universitaria como ecosistemas lingüísticos",
    "Examinar los espacios de enseñanza universitaria como entornos socioculturales y discursivos que influyen en la formación de competencias lingüísticas.",
    [
        "Learning spaces", "Educational spaces", "Academic discourse",
        "Communities of practice", "Language practices"
    ],
    [
        '("learning spaces" OR "communities of practice") AND ("academic discourse" OR "language practices") AND ("higher education")'
    ]
)

add_section(
    "T4. Formación lingüística situada y sociocultural",
    "Analizar la formación lingüística desde un enfoque situado y sociocultural, reconociendo la influencia del contexto institucional y social.",
    [
        "Situated learning", "Sociocultural approach",
        "Language socialization", "Academic literacy"
    ],
    [
        '("situated learning" AND "academic literacy") AND ("higher education")'
    ]
)

add_section(
    "T5. Las voces de docentes y discentes como categoría analítica",
    "Fundamentar el uso de las voces de docentes y discentes como categoría analítica para comprender la formación de competencias lingüísticas en la educación superior.",
    [
        "Student voice", "Teacher voice", "Educational discourse",
        "Narrative inquiry", "Interpretive approach"
    ],
    [
        '("student voice" OR "teacher voice") AND ("language education" OR "academic language")'
    ]
)

path = r"G:\Mi unidad\2025\Master FRANCISCO MARCELO ALVARADO PORRAS\Subtemas_Antecedentes_y_Marco_Teorico_Ecuaciones.docx"
doc.save(path)

path
